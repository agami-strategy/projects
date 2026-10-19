"""Label job ads with OSCA unit group codes (4-digit) using DeepSeek.

Approach: the full list of OSCA unit groups (code, title, lead statement, skill level,
occupation titles and alternative titles) goes in the system prompt, which DeepSeek caches
automatically by prefix. For each ad the model returns its top 3 candidate unit groups with
confidence and rationale, plus flags for ads that are not a single real occupation. Output is
a JSONL of raw results and an Excel review sheet with blank gold_code / reviewer_notes columns.

Usage:
    python scripts/label_osca.py data/samples/sample_100.jsonl --dry-run     # no API calls
    python scripts/label_osca.py data/samples/sample_100.jsonl               # label all
    python scripts/label_osca.py data/samples/sample_100.jsonl --limit 5     # try a few first
    python scripts/label_osca.py data/samples/sample_100.jsonl --model pro --fresh
    python scripts/label_osca.py data/samples/sample_100.jsonl --thinking    # slower, reasons before answering

Output files carry the model name: data/labels/osca_labels_<tag>_<model>.jsonl and osca_review_<tag>_<model>.xlsx.
Needs DEEPSEEK_API_KEY (environment or .env). Re-running resumes: ads already in the output JSONL are skipped.
"""
# %% Imports
# Where this script sits in the pipeline: sample_jobs.py produces a JSONL of ads (one JSON object
# per line with id, title, description_text, department, employment_type, locations, source_slug,
# apply_url). This script codes each ad to an OSCA 2024 unit group. label_seniority.py codes the
# same ads as entry-level vs experienced; the two outputs are joined on the ad id downstream to
# ask whether entry-level ads within an occupation describe more automatable tasks.
from __future__ import annotations

import argparse
import io
import json
import os
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Literal

import openpyxl
from pydantic import BaseModel, Field

# Make sibling modules importable when run as `python scripts/label_osca.py`. llm.py provides the
# DeepSeek call wrapper, the MODELS/price table and cost helpers; importing it also imports
# common.py, which installs the Excel control-character patch and the .env loader.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from llm import DEFAULT_MODEL, MODELS, USAGE_KEYS, Refusal, complete_json, cost, estimate_cost  # noqa: E402,F401

# %% Paths and constants
# ROOT = project folder (one level above scripts/). The OSCA descriptions workbook is the ABS
# release of OSCA 2024 with "Table 1" (occupations, 6-digit) and "Table 2" (unit groups, 4-digit).
ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OSCA = ROOT / "data" / "osca-descriptions.xlsx"
OUT_DIR = ROOT / "data" / "labels"
# Ad text is truncated to this many characters before it goes to the model. Design choice: it bounds
# the per-ad (uncached) input cost and long ads are mostly boilerplate after the first ~6k chars.
# If truncation were found to cut off requirements sections, raising this is the only change needed;
# the --dry-run cost estimate will reflect it because it measures the actual message length.
MAX_AD_CHARS = 6000  # ~1.5k tokens; responsibilities/qualifications are almost always within this


# --------------------------------------------------------------------------- OSCA loading
# %% Reading the OSCA workbook while Excel has it open (Windows only)
# Problem: a reviewer typically has osca-descriptions.xlsx open in Excel while running this script,
# and Excel opens the file with a lock that makes Python's normal open() raise PermissionError.
# Excel's lock does still permit readers that request full sharing (read + write + delete), which
# Python's open() does not do but the raw Win32 CreateFileW call can. The handle is converted to a
# C runtime file descriptor (msvcrt.open_osfhandle) and then to a normal Python file object, so the
# rest of the code just sees bytes. ctypes/msvcrt are imported inside the function because they are
# Windows-only and this path is only reached on win32.
def _read_bytes_shared(path: Path) -> bytes:
    """Read a file that Excel has open. Excel's lock rejects Python's normal open(), but a
    Win32 CreateFile with read+write+delete sharing is allowed."""
    import ctypes
    import msvcrt

    # Win32 constants: GENERIC_READ access; SHARE_ALL = FILE_SHARE_READ|WRITE|DELETE (0x1|0x2|0x4);
    # OPEN_EXISTING fails if the file is missing; FILE_ATTRIBUTE_NORMAL = no special flags.
    GENERIC_READ, SHARE_ALL, OPEN_EXISTING, FILE_ATTRIBUTE_NORMAL = 0x80000000, 0x7, 3, 0x80
    h = ctypes.windll.kernel32.CreateFileW(str(path), GENERIC_READ, SHARE_ALL, None, OPEN_EXISTING, FILE_ATTRIBUTE_NORMAL, None)
    # CreateFileW returns INVALID_HANDLE_VALUE (-1) on failure; GetLastError gives the Win32 code.
    if h == -1:
        raise PermissionError(f"cannot open {path} even with shared access (Win32 error {ctypes.GetLastError()})")
    # os.fdopen takes ownership of the descriptor (and thereby the handle), so the with-block closes it.
    with os.fdopen(msvcrt.open_osfhandle(h, os.O_RDONLY | os.O_BINARY), "rb") as f:
        return f.read()


# Normal path first; only on PermissionError and only on Windows fall back to the shared read and
# hand openpyxl an in-memory copy. read_only=True streams rows instead of loading cell objects,
# which is much faster for the ~1,200-row occupations table. On non-Windows the error propagates.
def load_workbook_safely(path: Path) -> openpyxl.Workbook:
    """openpyxl gets PermissionError while the workbook is open in Excel; fall back to a shared read."""
    try:
        return openpyxl.load_workbook(path, read_only=True)
    except PermissionError:
        if sys.platform != "win32":
            raise
        return openpyxl.load_workbook(io.BytesIO(_read_bytes_shared(path)), read_only=True)


# %% Parsing OSCA into {unit_group_code: {...}}
# Output: dict keyed by 4-digit code with title, lead statement, skill level and the list of
# 6-digit occupations (title, alternative titles, specialisations) nested under it. Expect 421 keys.
# Assumptions about the ABS workbook layout (worth re-checking if the ABS re-issues the file):
#   Table 2 columns: [0] code, [1] title, [2] lead statement, [4] skill level.
#   Table 1 columns: [0] 6-digit code, [1] title, [2] alternative titles, [9] specialisations.
# Rows are kept only when column 0 is purely digits, which drops headers, notes and blank rows.
# Codes are read as strings to preserve any leading zeros and to match the model's string output.
def load_osca(path: Path) -> dict[str, dict]:
    """Return {unit_group_code: {title, lead, skill_level, occupations:[{code,title,alt,spec}]}}."""
    wb = load_workbook_safely(path)
    t2 = [r for r in wb["Table 2"].iter_rows(values_only=True) if r[0] and str(r[0]).strip().isdigit()]
    t1 = [r for r in wb["Table 1"].iter_rows(values_only=True) if r[0] and str(r[0]).strip().isdigit()]

    groups: dict[str, dict] = {}
    for r in t2:
        code = str(r[0]).strip()
        # Table 2 also lists 1-, 2- and 3-digit aggregates (major/sub-major/minor groups); only the
        # 4-digit unit groups are wanted as coding targets.
        if len(code) != 4:
            continue
        groups[code] = {
            "code": code,
            "title": str(r[1]).strip(),
            "lead": str(r[2] or "").strip(),
            "skill_level": str(r[4] or "").strip(),
            "occupations": [],
        }
    for r in t1:
        code = str(r[0]).strip()
        # The first four digits of a 6-digit occupation code identify its unit group.
        ug = code[:4]
        if ug in groups:
            groups[ug]["occupations"].append({
                "code": code,
                "title": str(r[1] or "").strip(),
                "alt": str(r[2] or "").strip(),
                "spec": str(r[9] or "").strip(),
            })
    return groups


# %% One line of the prompt per unit group
# Formats a unit group as a single line for the system prompt: code, title, skill level, lead
# statement, then occupation titles and finally alternative titles + specialisations merged under
# "Also known as". Everything is on one line so the model can scan the 421-line list as a table.
# Design choice: including alt titles roughly doubles the listing size (and therefore the cached
# prefix) but is what lets foreign/inflated job titles map to Australian occupations.
def format_group(g: dict) -> str:
    occ_titles = "; ".join(o["title"] for o in g["occupations"] if o["title"])
    alts = "; ".join(x for o in g["occupations"] for x in (o["alt"], o["spec"]) if x)
    line = f"{g['code']} {g['title']} (skill level {g['skill_level']}): {g['lead']}"
    if occ_titles:
        line += f" Occupations: {occ_titles}."
    if alts:
        line += f" Also known as: {alts}."
    return line


# %% System prompt = coding instructions + full OSCA listing
# The prompt is built once per run and reused verbatim for every ad, so DeepSeek's automatic
# prefix cache serves it at the cache_hit rate (~2% of the miss price) after the first call. The
# whole thing is ~43k tokens, nearly all of it the listing; the per-ad user message is ~1-2k.
# Prompt design choices encoded in the instruction text below (a string literal; edit with care,
# any change invalidates the cache prefix for the next run but that is a one-off cost):
#   - code the occupation not the industry; code substantive duties over inflated titles;
#   - explicit OSCA skill-level definitions so the model can report ad_skill_level and use it to
#     choose between adjacent unit groups (this feeds the skill_mismatch flag in the review sheet);
#   - the dataset skew (tech/startup employers, many countries) is stated so foreign titles are
#     mapped rather than rejected;
#   - up to 3 candidates with confidences that need not sum to 1 (they are calibrated guesses,
#     not a distribution);
#   - an `issue` flag for ads that are not a single real job, so those can be excluded downstream;
#   - a worked JSON example, which in practice reduces malformed replies more than the schema alone.
# The listing is sorted by code so the prompt is deterministic across runs (stable cache prefix).
def build_system_prompt(groups: dict[str, dict]) -> str:
    instructions = """You are an expert occupation coder for the Australian Bureau of Statistics' OSCA \
(Occupation Standard Classification for Australia, 2024). You classify job advertisements to the \
4-digit OSCA unit group that best matches the work actually performed.

Coding rules:
- Code the occupation, not the industry or the employer. A software engineer at a bank is still a software engineer.
- Code the substantive work described in the responsibilities, not the title alone. Titles are inflated and vary by country: \
"Account Executive" in a US tech company is usually a sales representative, not a manager; a "Manager" who has no reports and \
does hands-on work is coded to the hands-on unit group; "Head of" / "Director" / "VP" with genuine control over a function is a manager.
- OSCA skill levels: 1 = bachelor degree or higher / 5+ years experience; 2 = diploma; 3 = certificate III/IV or trade; \
4 = certificate II; 5 = certificate I or short on-the-job training. Use the ad's stated requirements to judge the ad's skill level \
and let that inform the choice between adjacent unit groups.
- Most ads in this dataset come from technology, startup and knowledge-work employers in many countries. Map foreign titles to \
the closest Australian occupation.
- Give up to 3 candidates ordered best first. Confidence is your probability that the candidate is the correct unit group; \
the confidences do not need to sum to 1, but should sum to at most 1.
- Set issue when the ad is not a single, real, identifiable job: talent-community / expression-of-interest pages, ads covering \
several unrelated roles, placeholder or test postings, ads not in English (still attempt a code if you can read it), \
or ads with too little text to classify.

Answer in JSON. Example of the shape (values illustrative):
{"issue": "none", "ad_skill_level": 1, "candidates": [{"code": "2733", "title": "Software and Applications Engineers", \
"confidence": 0.85, "rationale": "designs and builds backend services"}, {"code": "2732", "title": "ICT Business and Systems \
Analysts and Architects", "confidence": 0.1, "rationale": "some architecture duties"}]}

The full OSCA unit group list follows. Return codes only from this list.

OSCA UNIT GROUPS
"""
    listing = "\n".join(format_group(g) for g in sorted(groups.values(), key=lambda g: g["code"]))
    return instructions + listing


# --------------------------------------------------------------------------- output schema
# %% Output schema (Pydantic)
# These classes serve two purposes: their JSON schema is appended to the system prompt by llm.py
# (so the model knows the field names and types) and the model's reply is validated against them
# locally. Field descriptions are visible to the model via the schema, so they double as prompt text.
# Note what is NOT enforced by the schema and therefore must be checked downstream:
#   - `code` is any string; validity against the 421 real codes is checked in write_review_xlsx
#     and counted in the run summary ("invalid codes returned");
#   - `confidence` is any float; the 0-1 range is only requested in the description;
#   - `ad_skill_level` is any int; 1-5 is only requested in the description;
#   - candidates length 1-3 is requested in the description; the review sheet slices to [:3].
# `issue` IS enforced as a Literal, so an unexpected issue string triggers a validation retry.
class Candidate(BaseModel):
    code: str = Field(description="4-digit OSCA unit group code")
    title: str = Field(description="unit group title exactly as listed")
    confidence: float = Field(description="probability between 0 and 1 that this is the correct unit group")
    rationale: str = Field(description="one sentence: which duties / requirements drove this choice")


class Label(BaseModel):
    issue: Literal["none", "talent_community", "multiple_roles", "not_a_job", "non_english", "insufficient_text"]
    ad_skill_level: int = Field(description="OSCA skill level 1-5 implied by the ad's stated requirements")
    candidates: list[Candidate] = Field(description="best first, 1 to 3 items; may be empty only if issue is not_a_job or insufficient_text")


# --------------------------------------------------------------------------- per-ad call
# %% Per-ad user message
# Builds the (uncached) user turn from one ad record. Included metadata: title, department,
# employment type, locations and the employer slug; these are cheap and often disambiguate
# (e.g. department "Sales" for an "Account Executive"). The employer slug is passed so the model
# can use industry context while the prompt tells it to code the occupation, not the industry.
# Truncation appends an explicit marker so the model knows text is missing rather than that the ad
# simply ends there. The closing sentence restates the task and the JSON requirement.
def ad_to_message(ad: dict) -> str:
    text = (ad.get("description_text") or "").strip()
    if len(text) > MAX_AD_CHARS:
        text = text[:MAX_AD_CHARS] + "\n[... truncated ...]"
    locs = ", ".join(ad.get("locations") or []) or "(none given)"
    return (
        f"JOB TITLE: {ad.get('title')}\n"
        f"DEPARTMENT: {ad.get('department') or '(none)'}\n"
        f"EMPLOYMENT TYPE: {ad.get('employment_type') or '(none)'}\n"
        f"LOCATION: {locs}\n"
        f"EMPLOYER SLUG: {ad.get('source_slug')}\n\n"
        f"AD TEXT:\n{text}\n\n"
        "Classify this ad to OSCA unit groups. Answer in JSON."
    )


# %% One labelled ad
# Wraps complete_json for a single ad and returns a JSONL-ready record. Two record shapes exist:
#   success: {"id", "label": {issue, ad_skill_level, candidates[]}, "usage": {...}, "model": "..."}
#   failure: {"id", "error": "..."}
# Presence of the "label" key is what every later step (resume, summary, review sheet) tests for.
# Refusal (no valid JSON after retries) is converted to an error record here; transport exceptions
# are left to propagate and are caught by work() in main(). max_tokens=2000 is ample for three
# candidates with one-sentence rationales; a reply hitting it would be retried as "cut off".
def label_one(system: str, ad: dict, model_key: str, thinking: bool) -> dict:
    try:
        parsed, usage, model_id = complete_json(system, ad_to_message(ad), Label, model_key, max_tokens=2000, thinking=thinking)
    except Refusal as e:
        return {"id": ad["id"], "error": f"no valid response: {e}"}
    return {"id": ad["id"], "label": parsed.model_dump(), "usage": usage, "model": model_id}


# --------------------------------------------------------------------------- review sheet
# %% Review workbook for human gold-labelling
# Produces osca_review_<tag>_<model>.xlsx with two sheets:
#   "review": one row per ad in input order, with ad metadata, the model's issue flag and skill
#             level, the top-3 candidates (code, resolved title, confidence, rationale), a skill
#             mismatch flag for the top candidate, and blank gold_code / gold_title / reviewer_notes
#             columns for the reviewer to fill in. Unlabelled ads get their error in the last column.
#   "osca_unit_groups": the full code list so the reviewer can look up codes without leaving Excel.
# The row layout is positional: the header has 26 columns and every branch below must append
# exactly 26 cells, so any change to `header` must be mirrored in each `row +=` line.
# Column widths, frozen panes (rows above 2 and columns before C stay visible) and the autofilter
# are cosmetic. Control characters in rationales are stripped by the patch in common.py.
def write_review_xlsx(ads: list[dict], results: dict[str, dict], groups: dict[str, dict], path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "review"
    header = [
        "id", "title", "company", "department", "location", "url",
        "issue", "ad_skill_level",
        "code_1", "title_1", "conf_1", "rationale_1", "ug_skill_1", "skill_mismatch_1",
        "code_2", "title_2", "conf_2", "rationale_2",
        "code_3", "title_3", "conf_3", "rationale_3",
        "gold_code", "gold_title", "reviewer_notes", "error",
    ]
    ws.append(header)
    for ad in ads:
        r = results.get(ad["id"])
        # First six cells: ad metadata straight from the input record.
        row = [
            ad["id"], ad.get("title"), ad.get("source_slug"), ad.get("department"),
            ", ".join(ad.get("locations") or []), ad.get("apply_url"),
        ]
        # Unlabelled ad (never attempted, or an error record): pad 16 label cells + 3 gold cells,
        # then put the error text in the final column. 6 + 16 + 3 + 1 = 26.
        if not r or "label" not in r:
            row += [None] * 16 + [None, None, None, (r or {}).get("error", "not labelled")]
            ws.append(row)
            continue
        lab = r["label"]
        row += [lab["issue"], lab["ad_skill_level"]]
        # Defensive slice: the schema does not cap the list length, the sheet has room for three.
        cands = lab["candidates"][:3]
        for i in range(3):
            if i < len(cands):
                c = cands[i]
                # The title is re-resolved from the OSCA table rather than trusting the model's title,
                # so a hallucinated code is visible as "INVALID CODE: <model title>" in the sheet.
                valid = c["code"] in groups
                title = groups[c["code"]]["title"] if valid else f"INVALID CODE: {c['title']}"
                row += [c["code"], title, round(c["confidence"], 2), c["rationale"]]
                if i == 0:
                    # Top candidate only: compare the model's ad_skill_level against the unit group's
                    # official skill level. OSCA lists some groups with several levels (e.g. "1; 2"),
                    # hence the split on ";" after removing spaces. "Y" marks a mismatch worth a look:
                    # either the ad is unusual for that group or the code is wrong.
                    ug_sl = groups[c["code"]]["skill_level"] if valid else ""
                    mismatch = valid and ug_sl and str(lab["ad_skill_level"]) not in ug_sl.replace(" ", "").split(";")
                    row += [ug_sl, "Y" if mismatch else ""]
            else:
                # Fewer than three candidates: pad the slot (plus the two skill cells if it is slot 1).
                row += [None, None, None, None] + ([None, None] if i == 0 else [])
        # gold_code, gold_title, reviewer_notes, error: all blank for a successfully labelled ad.
        row += [None, None, None, None]
        ws.append(row)

    # Second sheet: reference list of unit groups for the reviewer.
    ws2 = wb.create_sheet("osca_unit_groups")
    ws2.append(["code", "title", "skill_level", "lead_statement", "occupations"])
    for g in sorted(groups.values(), key=lambda g: g["code"]):
        ws2.append([g["code"], g["title"], g["skill_level"], g["lead"],
                    "; ".join(o["title"] for o in g["occupations"])])

    # Column letters are tied to the header order above (B=title, J/P/T=titles, L/R/V=rationales,
    # X=gold_title, Y=reviewer_notes). Reordering the header means re-mapping these.
    widths = {"B": 40, "C": 18, "E": 28, "J": 34, "L": 60, "P": 34, "R": 60, "T": 34, "V": 60, "X": 30, "Y": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# --------------------------------------------------------------------------- main
# %% CLI arguments
# Flags a reviewer is most likely to use:
#   --dry-run      build prompts, print sizes and a cost estimate, make no API calls (no key needed)
#   --limit N      label only the first N ads of the input (cheap smoke test)
#   --workers N    parallel API calls; default 8 (DeepSeek tolerates this; lower it on rate limits)
#   --model        flash (default, cheap) or pro; the model name is baked into the output file names
#   --thinking     DeepSeek reasoning mode; slower and more output tokens, sometimes better on hard ads
#   --fresh        delete the existing JSONL for this input+model and start again
#   --review-only  rebuild the xlsx from the JSONL without any API calls (e.g. after editing widths)
#   --tag          override the output-name suffix (default derives from the input file stem)
def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input", help="JSONL of ads from sample_jobs.py")
    ap.add_argument("--osca", default=str(DEFAULT_OSCA))
    ap.add_argument("--tag", default=None, help="output name suffix; default = input file stem")
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--model", default=DEFAULT_MODEL, choices=list(MODELS), help="short model name; output files carry it")
    ap.add_argument("--thinking", action="store_true", help="enable DeepSeek thinking mode (slower, more output tokens)")
    ap.add_argument("--dry-run", action="store_true", help="build prompts and report sizes; no API calls")
    ap.add_argument("--fresh", action="store_true", help="discard existing output for this input+model and relabel everything")
    ap.add_argument("--review-only", action="store_true", help="rebuild the review workbook from the existing JSONL; no API calls")
    args = ap.parse_args()

    # %% Load inputs: OSCA table, system prompt, ads
    # The ads file is read fully into memory (samples are hundreds to low thousands of ads).
    # --limit truncates AFTER loading, so the "ads" list and all counters refer to the limited set.
    groups = load_osca(Path(args.osca))
    system = build_system_prompt(groups)
    ads = [json.loads(l) for l in open(args.input, encoding="utf-8")]
    if args.limit:
        ads = ads[: args.limit]

    # %% Output paths
    # tag = <input stem without "sample_"> + "_" + <model key>, e.g. "100_australia_flash". Because the
    # model key is part of the name, flash and pro runs over the same input never overwrite or
    # resume from each other. Both the JSONL (raw, append-only) and the xlsx (regenerated) share it.
    tag = (args.tag or Path(args.input).stem.replace("sample_", "")) + "_" + args.model
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    jsonl_path = OUT_DIR / f"osca_labels_{tag}.jsonl"
    xlsx_path = OUT_DIR / f"osca_review_{tag}.xlsx"

    # %% Size report
    # The token estimate is the crude chars/4 rule; it is only used for the dry-run cost estimate and
    # the printed size. The real token counts come back in usage after each call.
    n_occ = sum(len(g["occupations"]) for g in groups.values())
    sys_tokens = len(system) // 4
    print(f"OSCA: {len(groups)} unit groups, {n_occ} occupations; system prompt ~{sys_tokens:,} tokens")
    print(f"ads: {len(ads)} from {args.input}; model {MODELS[args.model]['id']}")

    # %% --dry-run: show the first ad message and estimate cost, then stop
    # per_in adds 300 tokens to the first ad's message length as a rough allowance for the schema
    # text and message framing; 300 output tokens per ad is typical for three short candidates.
    # The first ad is used as the size proxy for all ads; a very short or long first ad skews it.
    if args.dry_run:
        print("\n--- first ad message ---\n" + ad_to_message(ads[0])[:1500] + "\n...")
        per_in = len(ad_to_message(ads[0])) // 4 + 300
        est = estimate_cost(len(ads), sys_tokens, per_in, 300, args.model)
        print(f"\nestimated cost at {MODELS[args.model]['id']} PEAK list price: ~${est:.2f} (off-peak is half); "
              f"assumes the ~{sys_tokens:,}-token OSCA prefix is served from cache after the first call")
        return

    # %% Resume logic
    # The JSONL is the source of truth for progress. Only records with a "label" key count as done;
    # error records are ignored on reload so a failed ad is retried on the next run. --fresh deletes
    # the file first. Because the file is opened in append mode below, a resumed run can contain
    # stale error lines from earlier runs followed by a later success for the same id; the review
    # sheet uses the in-memory `results` dict (last write wins per id), so that is harmless.
    results: dict[str, dict] = {}
    if args.fresh and jsonl_path.exists():
        jsonl_path.unlink()
        print("--fresh: discarded previous labels for this input and model")
    if jsonl_path.exists():
        for line in open(jsonl_path, encoding="utf-8"):
            r = json.loads(line)
            if "label" in r:
                results[r["id"]] = r
        print(f"resuming: {len(results)} already labelled")
    todo = [a for a in ads if a["id"] not in results]

    # %% --review-only: rebuild the workbook from existing labels and stop
    if args.review_only:
        write_review_xlsx(ads, results, groups, xlsx_path)
        print(f"wrote {xlsx_path} from {len(results)} existing labels ({len(todo)} ads unlabelled)")
        return

    # %% Labelling run: warm the cache, then fan out
    # The lock guards the shared `results` dict and the output file: worker threads do the API calls
    # but all writes happen on the main thread inside as_completed, so the lock is belt-and-braces
    # rather than strictly required. Each record is flushed immediately so a crash or Ctrl-C loses
    # at most the in-flight ads and the next run resumes from what was written.
    lock = threading.Lock()
    out = open(jsonl_path, "a", encoding="utf-8")
    t0 = time.time()

    # first call alone so the prefix cache is populated before the workers fan out
    # Why: DeepSeek caches prompt prefixes automatically, but only once a request with that prefix
    # has completed. If 8 workers started at once, all 8 would pay the full ~43k-token miss price.
    # Running one ad synchronously first costs a few seconds and makes every later call a hit.
    # The first result is printed with its usage so the reviewer can see cache_miss ~= prompt size.
    if todo:
        first = todo.pop(0)
        r = label_one(system, first, args.model, args.thinking)
        results[first["id"]] = r
        out.write(json.dumps(r, ensure_ascii=False) + "\n")
        out.flush()
        print(f"[1/{len(ads)}] {first['title'][:60]!r} -> "
              f"{[c['code'] for c in r.get('label', {}).get('candidates', [])] or r.get('error')}  usage={r.get('usage')}")

    # Worker wrapper: any exception that escapes label_one (network, SDK, unexpected response shape)
    # becomes an error record for that ad rather than killing the pool. Such ads are retried next run.
    def work(ad: dict) -> dict:
        try:
            return label_one(system, ad, args.model, args.thinking)
        except Exception as e:  # network / API errors are recorded, not fatal
            return {"id": ad["id"], "error": f"{type(e).__name__}: {e}"}

    # Submit everything, then consume in completion order (not input order); the progress counter
    # starts after the already-labelled count so [i/N] reflects overall progress on a resumed run.
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = {ex.submit(work, ad): ad for ad in todo}
        for i, fut in enumerate(as_completed(futs), start=len(results) + 1):
            ad = futs[fut]
            r = fut.result()
            with lock:
                results[ad["id"]] = r
                out.write(json.dumps(r, ensure_ascii=False) + "\n")
                out.flush()
            codes = [c["code"] for c in r.get("label", {}).get("candidates", [])] if "label" in r else r.get("error")
            print(f"[{i}/{len(ads)}] {ad['title'][:60]!r} -> {codes}")
    out.close()

    # %% Run summary
    # Totals cover everything in `results`, i.e. resumed labels plus this run's, so the cost line
    # on a resumed run overstates what THIS run spent. `invalid` counts candidate codes not in the
    # OSCA table (model hallucinations); a non-trivial rate would argue for a stricter prompt or a
    # post-hoc fuzzy match on title. The issues histogram shows how many ads were flagged as not a
    # single real job and should be excluded from the analysis.
    ok = [r for r in results.values() if "label" in r]
    errs = [r for r in results.values() if "label" not in r]
    usage = {k: sum(r["usage"][k] for r in ok) for k in USAGE_KEYS}
    invalid = sum(1 for r in ok for c in r["label"]["candidates"] if c["code"] not in groups)
    issues: dict[str, int] = {}
    for r in ok:
        issues[r["label"]["issue"]] = issues.get(r["label"]["issue"], 0) + 1
    print(f"\nlabelled {len(ok)}, errors {len(errs)}, invalid codes returned {invalid}, {time.time()-t0:.0f}s")
    print(f"tokens: {usage}  cost at peak rates ~${cost(usage, args.model):.2f} (off-peak half)")
    print(f"issues: {issues}")

    # %% Write the review workbook
    # Always regenerated from the full in-memory results, so it reflects resumed + new labels.
    write_review_xlsx(ads, results, groups, xlsx_path)
    print(f"wrote {jsonl_path}\nwrote {xlsx_path}")


if __name__ == "__main__":
    main()
