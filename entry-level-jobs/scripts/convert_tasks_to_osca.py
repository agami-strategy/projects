"""Convert ANZSCO unit-group task / time-share data to OSCA 2024 unit groups.

Inputs
    data/anzsco_task_share.xlsx           one row per (ANZSCO v1.3 unit group, task) with
                                          TOTAL_TIME_SHARE (sums to 1 per unit group) and its split
                                          into augmentation_time / automation_time / unaffected_time
    data/osca_correspondence_tables.xlsx  ABS correspondences; Table 1 = ANZSCO v1.3 -> OSCA 2024
                                          at 6-digit occupation level, 'p' = partial match, no proportions
    data/osca-descriptions.xlsx           OSCA titles and skill levels

Method
    The ABS correspondence carries no employment proportions, so weights are derived from the
    structure of the correspondence itself. For each OSCA unit group G, every ANZSCO unit group A
    contributes in proportion to the number of 6-digit occupation links A has into G (partial links
    count --partial-weight, default 1.0). A's task rows are copied into G with their time shares
    scaled by that weight, and rows sharing a Task ID are summed. Time shares therefore still sum
    to 1 within each OSCA unit group. Task-level attributes (exposure labels, predictions) are
    unchanged because they describe the task, not the occupation.

    A 1:1 mapping reproduces the ANZSCO row exactly. A split (one ANZSCO group feeding several
    OSCA groups) gives each OSCA group the full ANZSCO task list. A merge (several ANZSCO groups
    feeding one OSCA group) blends their task lists by link count.

Usage
    python scripts/convert_tasks_to_osca.py
    python scripts/convert_tasks_to_osca.py --partial-weight 0.5

Output  data/osca_task_share.xlsx (sheets: tasks, summary, mapping, unmapped) and osca_task_share.csv
"""
# %% Imports and file paths
# This is the data-preparation step that produces data/osca_task_share.csv, the task file every
# downstream script (match_tasks, regress_*, task_prevalence) reads. The task/time-share data exist
# only for ANZSCO v1.3 unit groups, while the job ads are coded to OSCA 2024, so a crosswalk is needed.
# Everything is done with openpyxl and plain dicts; no pandas.
from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

# label_osca supplies: DEFAULT_OSCA (path to the OSCA descriptions workbook), load_osca (code -> title,
# skill level) and load_workbook_safely (openpyxl loader that tolerates the ABS files' quirks).
sys.path.insert(0, str(Path(__file__).resolve().parent))
from label_osca import DEFAULT_OSCA, load_osca, load_workbook_safely  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
TASK_FILE = ROOT / "data" / "anzsco_task_share.xlsx"
CORR_FILE = ROOT / "data" / "osca_correspondence_tables.xlsx"
OUT_XLSX = ROOT / "data" / "osca_task_share.xlsx"
OUT_CSV = ROOT / "data" / "osca_task_share.csv"

# %% Column conventions of the source task workbook
# TIME_COLS are the four per-task time quantities that get rescaled by the crosswalk weight and summed
# across sources. Note "augmetnation_time" is the source workbook's own misspelled column name; it is
# read verbatim and renamed to augmentation_time only in the output header below.
# NUMERIC_TASK_COLS are model predictions (continuous) that are averaged across sources, weighted by
# time share. LABEL_TASK_COLS are categorical/dummy task attributes copied from the first source seen,
# on the assumption that a given Task ID carries the same labels wherever it appears.
CORR_SHEET = "Table 1"  # ANZSCO v1.3 -> OSCA 2024 v1.0 (the task file uses v1.3 unit groups)
TIME_COLS = ["TOTAL_TIME_SHARE", "augmetnation_time", "automation_time", "unaffected_time"]
NUMERIC_TASK_COLS = ["aug_pred", "ai_pred"]
LABEL_TASK_COLS = ["gpt4_exposure", "gpt4_automation", "ai_dummy", "aug_dummy"]


# %% Load the ANZSCO task rows
# Reads the first worksheet; the first row is the header. A blank in column B (index 1) marks an
# empty/trailing row and is skipped. Codes are stringified and stripped so that a numeric cell
# (e.g. 2211 stored as int) and a text cell ("2211") both key the same group.
# Returns titles {anzsco_ug: title} and tasks {anzsco_ug: [row dict keyed by header]}.
def load_tasks(path: Path) -> tuple[dict[str, str], dict[str, list[dict]]]:
    """Return ({anzsco_ug: title}, {anzsco_ug: [row dicts]})."""
    ws = load_workbook_safely(path).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "_" for h in next(rows)]
    titles: dict[str, str] = {}
    tasks: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        if r[1] is None:
            continue
        d = dict(zip(header, r))
        code = str(d["ANZSCO_UNIT_CODE"]).strip()
        titles[code] = str(d["ANZSCO_TITLE"]).strip()
        tasks[code].append(d)
    return titles, tasks


# %% Load the ABS correspondence (6-digit occupation links)
# Layout of "Table 1" in the ABS workbook: five title/header rows (skipped by [5:]), then one row per
# link with column A = ANZSCO v1.3 6-digit code, column C = OSCA 6-digit code, column D = "p" for a
# partial match. When one ANZSCO occupation maps to several OSCA occupations the ANZSCO code is shown
# once and left blank on the following rows (visually merged cells), hence `src` carries forward.
# Rows whose column C is not all digits (notes, blank lines) are ignored.
# Output: a flat list of (anzsco6, osca6, is_partial) tuples; the 4-digit unit group is derived later.
def load_correspondence(path: Path, sheet: str) -> list[tuple[str, str, bool]]:
    """Return [(anzsco6, osca6, is_partial)]. Blank source cells continue the previous source."""
    ws = load_workbook_safely(path)[sheet]
    out, src = [], None
    for r in list(ws.iter_rows(values_only=True))[5:]:
        if r[0]:
            src = str(r[0]).strip()
        if src and r[2] and str(r[2]).strip().isdigit():
            out.append((src, str(r[2]).strip(), str(r[3] or "").strip().lower() == "p"))
    return out


# %% Crosswalk weights from link counts
# The core modelling choice. The ABS correspondence has no employment or proportion data, so the only
# information about "how much" of ANZSCO group A belongs in OSCA group G is the number of 6-digit
# occupations linking A to G. The weight of A in G is
#     score(A, G) / sum over sources A' with task data of score(A', G)
# where score counts full links as 1 and partial links as --partial-weight. Example: if G is fed by
# A (3 links) and B (1 link), A gets 0.75 and B 0.25. If B has no task data, A gets 1.0 and B is kept
# with weight 0 so the gap is visible in the mapping sheet. Codes are cut to 4 digits ([:4]) because
# both classifications use the first four digits as the unit group.
# What would change: with employment weights (if they existed) large occupations would dominate merges;
# with --partial-weight 0.5 partial links count half, shifting weight towards full matches.
def build_weights(links: list[tuple[str, str, bool]], partial_weight: float, has_tasks: set[str]) -> dict[str, dict[str, dict]]:
    """{osca_ug: {anzsco_ug: {links, partial, weight}}} with weights summing to 1 per OSCA unit group.

    Weights are normalised over source ANZSCO groups that have task data; a source without tasks
    is kept in the table with weight 0 so the gap is visible in the mapping sheet.
    """
    # raw[osca4][anzsco4] accumulates link counts; the nested defaultdict creates cells on first access
    raw: dict[str, dict[str, dict]] = defaultdict(lambda: defaultdict(lambda: {"links": 0, "partial": 0, "score": 0.0}))
    for a6, o6, partial in links:
        cell = raw[o6[:4]][a6[:4]]
        cell["links"] += 1
        cell["partial"] += int(partial)
        cell["score"] += partial_weight if partial else 1.0
    weights: dict[str, dict[str, dict]] = {}
    for og, srcs in raw.items():
        # denominator restricted to sources that actually have task rows, so weights over those sum to 1
        total = sum(c["score"] for ag, c in srcs.items() if ag in has_tasks)
        weights[og] = {ag: {"links": c["links"], "partial": c["partial"],
                            "weight": c["score"] / total if (total and ag in has_tasks) else 0.0}
                       for ag, c in srcs.items()}
    return weights


# %% Apply the weights: build each OSCA group's task list
# For each OSCA group, copy every task row of every weighted source, scaling the four time columns by
# the source's weight, and merge rows that share a Task ID by summing. Because each ANZSCO group's
# TOTAL_TIME_SHARE sums to 1 and the weights sum to 1, the OSCA group's shares sum to 1 as well
# (checked by the assertion in main). The three components stay consistent with the total because
# they are scaled by the same factor.
# Predictions (aug_pred, ai_pred) are combined as a time-share-weighted mean: accumulate pred * share
# in _pred_acc, then divide by the merged share. For a task that appears in only one source this is
# just the original prediction. Label columns are taken from the first source that contributes the task.
def convert(tasks: dict[str, list[dict]], weights: dict[str, dict[str, dict]]) -> dict[str, list[dict]]:
    """Return {osca_ug: [task rows with rescaled time shares]}."""
    out: dict[str, list[dict]] = {}
    for og, srcs in weights.items():
        merged: dict[str, dict] = {}
        for ag, w in srcs.items():
            if w["weight"] == 0 or ag not in tasks:
                continue
            for t in tasks[ag]:
                tid = str(t["Task ID"])
                # this source's contribution to the task's OSCA time share
                share = float(t["TOTAL_TIME_SHARE"]) * w["weight"]
                if tid not in merged:
                    merged[tid] = {
                        "Task ID": tid, "Task": t["Task"],
                        **{c: t.get(c) for c in LABEL_TASK_COLS},
                        **{c: 0.0 for c in TIME_COLS},
                        "_pred_acc": {c: 0.0 for c in NUMERIC_TASK_COLS},
                        "source_anzsco": [],
                    }
                m = merged[tid]
                for c in TIME_COLS:
                    m[c] += float(t[c]) * w["weight"]
                for c in NUMERIC_TASK_COLS:
                    if t.get(c) is not None:
                        m["_pred_acc"][c] += float(t[c]) * share
                m["source_anzsco"].append(ag)
        rows = []
        for m in merged.values():
            share = m["TOTAL_TIME_SHARE"]
            for c in NUMERIC_TASK_COLS:  # share-weighted mean of predictions across source groups
                m[c] = m["_pred_acc"][c] / share if share else None
            del m["_pred_acc"]
            # provenance: which ANZSCO groups contributed this task, as a ";"-separated string
            m["source_anzsco"] = ";".join(sorted(set(m["source_anzsco"])))
            rows.append(m)
        # largest time share first, so the most time-consuming tasks head each group's list
        rows.sort(key=lambda m: -m["TOTAL_TIME_SHARE"])
        out[og] = rows
    return out


# %% Command-line entry point
def main() -> None:
    # %% Arguments and inputs
    # All three input paths default to the data/ files listed in the module docstring. --partial-weight
    # is the only tuning knob; the default 1.0 treats partial and full links alike.
    ap = argparse.ArgumentParser()
    ap.add_argument("--tasks", default=str(TASK_FILE))
    ap.add_argument("--correspondence", default=str(CORR_FILE))
    ap.add_argument("--osca", default=str(DEFAULT_OSCA))
    ap.add_argument("--partial-weight", type=float, default=1.0,
                    help="weight of a partial ('p') occupation link relative to a full one (default 1.0)")
    args = ap.parse_args()

    # set(tasks) is the set of ANZSCO groups that have task data, used to normalise the weights
    anzsco_titles, tasks = load_tasks(Path(args.tasks))
    links = load_correspondence(Path(args.correspondence), CORR_SHEET)
    groups = load_osca(Path(args.osca))
    weights = build_weights(links, args.partial_weight, set(tasks))
    converted = convert(tasks, weights)

    # %% Diagnostics: mapping type, gaps, and the sum-to-one check
    # mapping_type per OSCA group: n_src counts contributing sources (with tasks); split is true if any
    # source also feeds another OSCA group. "1:1" reproduces an ANZSCO row exactly; "split" copies a
    # full task list into several OSCA groups; "merge" blends several lists; "merge+split" both.
    # osca_no_tasks: OSCA groups in the descriptions file that end up with no task rows (no correspondence
    # or only sources without tasks). anzsco_unmapped: task groups the correspondence never mentions.
    # The assertion is the invariant the method promises: time shares sum to 1 in every populated group.
    # It fails if a source ANZSCO group's shares do not sum to 1 in the input workbook.
    # provenance / diagnostics
    mapping_type = {}
    for og, srcs in weights.items():
        n_src = sum(1 for ag in srcs if ag in tasks)
        split = any(len([g for g, s in weights.items() if ag in s]) > 1 for ag in srcs)
        mapping_type[og] = ("merge+split" if n_src > 1 and split else "merge" if n_src > 1 else "split" if split else "1:1")
    osca_no_tasks = sorted(set(groups) - {og for og, rows in converted.items() if rows})
    anzsco_unmapped = sorted(set(tasks) - {ag for srcs in weights.values() for ag in srcs})
    for og, rows in converted.items():
        s = sum(r["TOTAL_TIME_SHARE"] for r in rows)
        if rows and abs(s - 1) > 1e-6:
            raise AssertionError(f"time shares in OSCA {og} sum to {s}")

    # %% Assemble output rows
    # tasks sheet / CSV: one row per OSCA group x task. The header renames TOTAL_TIME_SHARE to TIME_SHARE
    # and the misspelled source column to augmentation_time; downstream scripts read these names.
    # Groups missing from the descriptions file get title "?" rather than failing.
    # ---- write
    task_header = ["OSCA_UNIT_CODE", "OSCA_TITLE", "Task ID", "Task", *LABEL_TASK_COLS, *NUMERIC_TASK_COLS,
                   "TIME_SHARE", "augmentation_time", "automation_time", "unaffected_time", "source_anzsco"]
    task_rows = []
    for og in sorted(converted):
        for r in converted[og]:
            task_rows.append([og, groups.get(og, {}).get("title", "?"), r["Task ID"], r["Task"],
                              *(r[c] for c in LABEL_TASK_COLS), *(r[c] for c in NUMERIC_TASK_COLS),
                              r["TOTAL_TIME_SHARE"], r["augmetnation_time"], r["automation_time"], r["unaffected_time"],
                              r["source_anzsco"]])

    # summary sheet: one row per OSCA group (including those with no tasks), with the mapping type, the
    # sources and their weights as text, and the group's total augmentation / automation / unaffected time
    # (these three sum to 1 for a populated group and are the occupation-level exposure baseline).
    summary_header = ["OSCA_UNIT_CODE", "OSCA_TITLE", "skill_level", "mapping_type", "n_tasks", "n_source_anzsco",
                      "source_anzsco_weights", "augmentation_time", "automation_time", "unaffected_time"]
    summary_rows = []
    for og in sorted(groups):
        rows = converted.get(og, [])
        srcs = weights.get(og, {})
        summary_rows.append([
            og, groups[og]["title"], groups[og]["skill_level"], mapping_type.get(og, "no correspondence"),
            len(rows), sum(1 for ag in srcs if ag in tasks),
            "; ".join(f"{ag} {anzsco_titles.get(ag, '?')} ({w['weight']:.2f})" for ag, w in sorted(srcs.items(), key=lambda kv: -kv[1]["weight"])),
            sum(r["augmetnation_time"] for r in rows) if rows else None,
            sum(r["automation_time"] for r in rows) if rows else None,
            sum(r["unaffected_time"] for r in rows) if rows else None,
        ])

    # mapping sheet: one row per (OSCA group, ANZSCO source) pair with link counts and the weight, so a
    # reviewer can audit every crosswalk decision; anzsco_has_tasks shows zero-weight sources
    mapping_rows = [[og, groups.get(og, {}).get("title", "?"), ag, anzsco_titles.get(ag, "(no tasks)"),
                     w["links"], w["partial"], round(w["weight"], 4), ag in tasks]
                    for og, srcs in sorted(weights.items()) for ag, w in sorted(srcs.items())]

    # %% Write the workbook and the CSV
    # openpyxl: ws.append writes one row at a time; freeze_panes "C2" keeps the header row and the first
    # two columns (code, title) visible while scrolling; auto_filter.ref = ws.dimensions puts filter
    # dropdowns on the whole used range. The CSV duplicates the tasks sheet for scripts that read CSV;
    # utf-8-sig writes a BOM so Excel detects the encoding, newline="" prevents blank lines on Windows.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "tasks"
    ws.append(task_header)
    for r in task_rows:
        ws.append(r)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws2 = wb.create_sheet("summary")
    ws2.append(summary_header)
    for r in summary_rows:
        ws2.append(r)
    ws2.freeze_panes = "C2"
    ws2.auto_filter.ref = ws2.dimensions
    ws3 = wb.create_sheet("mapping")
    ws3.append(["OSCA_UNIT_CODE", "OSCA_TITLE", "ANZSCO_UNIT_CODE", "ANZSCO_TITLE", "occupation_links", "partial_links", "weight", "anzsco_has_tasks"])
    for r in mapping_rows:
        ws3.append(r)
    ws4 = wb.create_sheet("unmapped")
    ws4.append(["issue", "code", "title"])
    for og in osca_no_tasks:
        ws4.append(["OSCA unit group with no task data", og, groups[og]["title"]])
    for ag in anzsco_unmapped:
        ws4.append(["ANZSCO unit group with tasks but no OSCA correspondence", ag, anzsco_titles[ag]])
    wb.save(OUT_XLSX)
    with open(OUT_CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(task_header)
        w.writerows(task_rows)

    # %% Console summary
    # Row counts to sanity-check: OSCA task rows exceed ANZSCO rows when splits outnumber merges.
    print(f"ANZSCO unit groups with tasks: {len(tasks)} | correspondence links: {len(links)} "
          f"({sum(1 for l in links if l[2])} partial) | OSCA unit groups: {len(groups)}")
    print(f"OSCA unit groups with task data: {sum(1 for og in converted if converted[og])} | mapping types: {dict(Counter(mapping_type.values()))}")
    print(f"OSCA unit groups without task data: {len(osca_no_tasks)} | ANZSCO groups with tasks but no correspondence: {len(anzsco_unmapped)}")
    print(f"task rows: ANZSCO {sum(len(v) for v in tasks.values()):,} -> OSCA {len(task_rows):,}")
    print(f"wrote {OUT_XLSX}\nwrote {OUT_CSV}")


if __name__ == "__main__":
    main()
