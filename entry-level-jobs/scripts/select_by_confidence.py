"""Per OSCA unit group, pick the highest OSCA-confidence cutoff that still leaves at least
--min-ads ads (falling back to --floor when none does), and optionally write the selected ads
as a new sample file for the seniority / task stages.

Rule: for each group, try the thresholds from highest to lowest; the first with >= min-ads
ads is chosen. If none qualifies, the floor threshold is used and all ads at or above it are
kept. Only ads with no issue flag and a valid top-1 code are considered.

Usage:
    python scripts/select_by_confidence.py data/samples/sample_us_2026-09-17.jsonl                # table only
    python scripts/select_by_confidence.py data/samples/sample_us_2026-09-17.jsonl --write        # + sample file
    python scripts/select_by_confidence.py ... --target data/osca_anzsco_target.xlsx --u25-only  # restrict groups
    python scripts/select_by_confidence.py ... --min-ads 1000 --thresholds 0.9 0.8 0.7 --floor 0.5 --cap 1500

Outputs:
    data/selection_<tag>_<model>.xlsx                     the table (one row per group)
    with --write: data/samples/sample_<tag>_selected.jsonl and data/labels/osca_labels_<tag>_selected_<model>.jsonl
"""
# %% Imports
# Standard library plus openpyxl (Excel output). `llm` and `label_osca` are sibling modules in scripts/:
#   MODELS / DEFAULT_MODEL   registry of model keys -> ids and prices (used only to build file names here)
#   load_osca                {osca_code: {"title", "skill_level", ...}} from the OSCA 2024 list
#   load_workbook_safely     openpyxl loader that tolerates the ABS workbook's quirks
from __future__ import annotations

import argparse
import json
import random
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

# Make the sibling modules importable when run as `python scripts/select_by_confidence.py`.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from llm import DEFAULT_MODEL, MODELS  # noqa: E402
from label_osca import DEFAULT_OSCA, load_osca, load_workbook_safely  # noqa: E402

# %% Directory layout
# ROOT is the project directory (one level above scripts/). OSCA label files live in data/labels and
# sample JSONLs in data/samples; every file name is built from the input "tag" derived in main().
ROOT = Path(__file__).resolve().parents[1]
LABEL_DIR = ROOT / "data" / "labels"
SAMPLE_DIR = ROOT / "data" / "samples"


# %% Optional restriction to youth-relevant occupations (--u25-only)
# data/osca_anzsco_target.xlsx has two sheets:
#   "Sheet1"  one row per ANZSCO unit group: column A = ANZSCO code, column D = U25 flag ("1" means the
#             occupation is a target for under-25 workers).
#   "Table 1" the ABS ANZSCO -> OSCA correspondence: column B = ANZSCO code (blank on continuation rows
#             when one ANZSCO unit maps to several OSCA groups), column F = OSCA unit code. The first
#             5 rows are titles/headers and are skipped.
# Returns the set of OSCA codes reachable from any U25 = 1 ANZSCO unit. A reviewer should check the
# column positions (r[0], r[3], r[1], r[5]) and the 5-row header skip against the current workbook;
# if the ABS changes the layout this function silently returns the wrong set.
def u25_osca_groups(target_path: Path) -> set[str]:
    """OSCA unit groups linked (via the workbook's Table 1) to any ANZSCO unit with U25 = 1."""
    wb = load_workbook_safely(target_path)
    # ANZSCO codes flagged U25 = 1. [1:] skips the header row; the flag is compared as a string.
    targets = {str(r[0]).strip() for r in list(wb["Sheet1"].iter_rows(values_only=True))[1:] if r[0] and str(r[3]) == "1"}
    # Walk the correspondence table. `src` carries the last seen ANZSCO code forward across merged-cell
    # continuation rows, so every OSCA row under a U25 ANZSCO unit is picked up.
    out, src = set(), None
    for r in list(wb["Table 1"].iter_rows(values_only=True))[5:]:
        if r[1]:
            src = str(r[1]).strip()
        if src in targets and r[5]:
            out.add(str(r[5]).strip())
    return out


# %% Command line
# input         the ads JSONL that label_osca.py was run on; its stem gives the tag (e.g. "us_2026-09-17")
# --model       which model's OSCA label file to read (osca_labels_<tag>_<model>.jsonl)
# --min-ads     a cutoff qualifies for a group only if >= this many ads clear it (default 1,000)
# --thresholds  the ladder of confidence cutoffs tried from highest to lowest (default 0.9 0.8 0.7)
# --floor       fallback cutoff when no ladder step reaches --min-ads (default 0.5)
# --cap         optional seeded random cap on ads per group, applied after the cutoff
# --u25-only    restrict to OSCA groups linked to U25 = 1 ANZSCO units (see u25_osca_groups)
# --write       also emit the selected ads + their OSCA labels as a new "<tag>_selected" sample
def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input", help="ads JSONL that label_osca.py was run on")
    ap.add_argument("--model", default=DEFAULT_MODEL, choices=list(MODELS))
    ap.add_argument("--min-ads", type=int, default=1000)
    ap.add_argument("--thresholds", type=float, nargs="*", default=[0.9, 0.8, 0.7])
    ap.add_argument("--floor", type=float, default=0.5)
    ap.add_argument("--cap", type=int, default=None, help="random cap per group after selection (seeded)")
    ap.add_argument("--target", default=str(ROOT / "data" / "osca_anzsco_target.xlsx"))
    ap.add_argument("--u25-only", action="store_true", help="only groups linked to U25 = 1 ANZSCO units")
    ap.add_argument("--write", action="store_true", help="write the selected ads + their OSCA labels as a new sample")
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()

    # %% Tag, OSCA list and the allowed set of groups
    # tag: "sample_us_2026-09-17.jsonl" -> "us_2026-09-17"; all input/output names derive from it.
    # allowed: the OSCA codes a top-1 label may take to be considered. Without --u25-only it is every
    # code in the OSCA list, so the U25 restriction is the only thing that narrows the occupation set.
    tag = Path(args.input).stem.replace("sample_", "")
    groups = load_osca(Path(DEFAULT_OSCA))
    allowed = u25_osca_groups(Path(args.target)) if args.u25_only else set(groups)

    # %% Load the eligible OSCA labels
    # One record per ad: {"id", "label": {"issue", "candidates": [{"code", "confidence"}, ...]}}.
    # An ad is eligible only if it has a label, no issue flag ("none"), at least one candidate, and its
    # top-1 code is in `allowed`. Ads failing any check never count toward a group's size below.
    # Assumption: candidates[0] is the model's top choice (label_osca.py writes them in rank order).
    labels: dict[str, dict] = {}
    for line in open(LABEL_DIR / f"osca_labels_{tag}_{args.model}.jsonl", encoding="utf-8"):
        r = json.loads(line)
        lab = r.get("label")
        if lab and lab["issue"] == "none" and lab["candidates"] and lab["candidates"][0]["code"] in allowed:
            labels[r["id"]] = r

    # %% Group the eligible ads by top-1 OSCA code
    # by_group[code] = [(top-1 confidence, ad id), ...]. Only the top-1 candidate matters from here on;
    # the 2nd/3rd candidates from label_osca.py are not used by this script.
    by_group: dict[str, list[tuple[float, str]]] = defaultdict(list)  # code -> [(conf, id)]
    for id_, r in labels.items():
        c = r["label"]["candidates"][0]
        by_group[c["code"]].append((c["confidence"], id_))

    # %% Choose a confidence cutoff per group
    # ladder = the --thresholds list, de-duplicated and sorted high -> low. For every allowed group:
    #   counts[t]  number of ads with top-1 confidence >= t, for each ladder step and for the floor
    #   chosen     the first (highest) ladder step whose count reaches --min-ads; None if none does
    #   rule       "min-ads met" or "floor", recorded in the table so the choice is auditable
    # If the ladder fails, the floor is used and ALL ads at or above it are kept: small groups keep
    # whatever they have at conf >= floor rather than being dropped, so the floor is the effective
    # minimum label quality anywhere in the sample.
    # --cap, if set, trims a group at random using one seeded RNG shared across groups, so the draw
    # depends on --seed and on the (sorted) group order. Groups with no eligible ads still get a row.
    # Tweaks: a higher --min-ads pushes more groups to lower cutoffs; a lower --floor admits noisier
    # labels; changing the ladder changes only which cutoffs are tried, not the rule.
    ladder = sorted(set(args.thresholds), reverse=True)
    rows, selected_ids = [], set()
    rng = random.Random(args.seed)
    for code in sorted(allowed):
        items = by_group.get(code, [])
        # ads clearing each cutoff; Excel equivalent: COUNTIFS(code, g, confidence, ">=" & t)
        counts = {t: sum(1 for conf, _ in items if conf >= t) for t in ladder + [args.floor]}
        # highest ladder step meeting --min-ads, else None -> fall back to the floor
        chosen = next((t for t in ladder if counts[t] >= args.min_ads), None)
        rule = "min-ads met" if chosen is not None else "floor"
        chosen = chosen if chosen is not None else args.floor
        ids = [i for conf, i in items if conf >= chosen]
        # seeded random cap, only when the group exceeds it (sampling without replacement)
        if args.cap and len(ids) > args.cap:
            ids = rng.sample(ids, args.cap)
        selected_ids.update(ids)
        rows.append([code, groups[code]["title"], len(items), *(counts[t] for t in ladder), counts[args.floor],
                     chosen, rule, len(ids)])

    # %% Write the selection table (one row per allowed group)
    # Columns: code, title, ads at any confidence, ads per ladder step, ads at the floor, cutoff used,
    # rule, ads selected. Sorted by ads_selected descending; header frozen and autofilter on.
    # This table is the audit trail for which cutoff each occupation got.
    header = ["osca_code", "osca_title", "ads_any_conf", *(f"ads_conf>={t}" for t in ladder), f"ads_conf>={args.floor}",
              "threshold_used", "rule", "ads_selected"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "selection"
    ws.append(header)
    for r in sorted(rows, key=lambda r: -r[-1]):
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out_xlsx = ROOT / "data" / f"selection_{tag}_{args.model}.xlsx"
    wb.save(out_xlsx)

    # %% Console summary
    # Row positions: r[-1] = ads_selected, r[-3] = threshold_used. "eligible" = ads that passed the
    # label filter above, before any cutoff.
    n_groups = sum(1 for r in rows if r[-1] > 0)
    print(f"groups considered: {len(rows)} | with ads: {n_groups} | rule: highest of {ladder} with >= {args.min_ads} ads, else floor {args.floor}")
    print(f"threshold used: {dict(Counter(r[-3] for r in rows if r[-1] > 0))} | ads selected: {len(selected_ids):,} of {len(labels):,} eligible")
    print(f"wrote {out_xlsx}")

    # %% Optionally write the selected sample (--write)
    # Re-reads the input JSONL so ads are written in their original order and copied byte-for-byte
    # (the raw line is written, not re-serialised). The OSCA label file is written in the same order
    # with the same ids under the tag "<tag>_selected", so label_seniority.py and sample_for_tasks.py
    # find both files by tag without any extra arguments.
    if args.write:
        out_sample = SAMPLE_DIR / f"sample_{tag}_selected.jsonl"
        out_labels = LABEL_DIR / f"osca_labels_{tag}_selected_{args.model}.jsonl"
        n = 0
        with open(out_sample, "w", encoding="utf-8") as fs, open(out_labels, "w", encoding="utf-8") as fl:
            for line in open(args.input, encoding="utf-8"):
                ad = json.loads(line)
                if ad["id"] in selected_ids:
                    fs.write(line if line.endswith("\n") else line + "\n")
                    fl.write(json.dumps(labels[ad["id"]], ensure_ascii=False) + "\n")
                    n += 1
        print(f"wrote {n:,} ads -> {out_sample}\nwrote matching OSCA labels -> {out_labels}")
        print(f"next: python scripts/label_seniority.py {out_sample.relative_to(ROOT).as_posix()}")


if __name__ == "__main__":
    main()
