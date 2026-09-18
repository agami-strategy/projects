"""Label job ads with OSCA unit group codes (4-digit) using Claude.

Approach: the full list of OSCA unit groups (code, title, lead statement, skill level,
occupation titles and alternative titles) goes in a cached system prompt. For each ad
Claude returns its top 3 candidate unit groups with confidence and rationale, plus flags
for ads that are not a single real occupation. Output is a JSONL of raw results and an
Excel review sheet with blank gold_code / reviewer_notes columns for manual correction.

Usage:
    python scripts/label_osca.py data/samples/sample_100.jsonl --dry-run     # no API calls
    python scripts/label_osca.py data/samples/sample_100.jsonl               # label all
    python scripts/label_osca.py data/samples/sample_100.jsonl --limit 5     # try a few first
    python scripts/label_osca.py data/samples/sample_100.jsonl --limit 5 --model haiku-4-5 --fresh   # relabel, ignore existing output
    python scripts/label_osca.py data/samples/sample_100_australia.jsonl --tag aus

Output files carry the model name: data/labels/osca_labels_<tag>_<model>.jsonl and osca_review_<tag>_<model>.xlsx.

Needs ANTHROPIC_API_KEY, either as an environment variable or in entry-level-work/.env
(copy .env.example). Re-running resumes: ads already present
in the output JSONL are skipped.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Literal

import anthropic
import openpyxl
from pydantic import BaseModel, Field

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import load_env  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OSCA = ROOT / "data" / "osca-descriptions.xlsx"
OUT_DIR = ROOT / "data" / "labels"

# List prices in USD per million tokens (cache write = 1.25x input, cache read = 0.1x input).
# `effort` is only accepted by the Claude 4.6+ generation; Haiku 4.5 rejects it.
MODELS = {
    "opus-5":   {"id": "claude-opus-5",   "price": {"input": 5.0, "cache_write": 6.25, "cache_read": 0.50, "output": 25.0}, "effort": True},
    "sonnet-5": {"id": "claude-sonnet-5", "price": {"input": 2.0, "cache_write": 2.50, "cache_read": 0.20, "output": 10.0}, "effort": True},
    "haiku-4-5": {"id": "claude-haiku-4-5", "price": {"input": 1.0, "cache_write": 1.25, "cache_read": 0.10, "output": 5.0}, "effort": False},
}
DEFAULT_MODEL = "haiku-4-5"

MAX_AD_CHARS = 6000  # ~1.5k tokens; responsibilities/qualifications are almost always within this


# --------------------------------------------------------------------------- OSCA loading
def _read_bytes_shared(path: Path) -> bytes:
    """Read a file that Excel has open. Excel's lock rejects Python's normal open(), but a
    Win32 CreateFile with read+write+delete sharing is allowed."""
    import ctypes
    import msvcrt

    GENERIC_READ, SHARE_ALL, OPEN_EXISTING, FILE_ATTRIBUTE_NORMAL = 0x80000000, 0x7, 3, 0x80
    h = ctypes.windll.kernel32.CreateFileW(str(path), GENERIC_READ, SHARE_ALL, None, OPEN_EXISTING, FILE_ATTRIBUTE_NORMAL, None)
    if h == -1:
        raise PermissionError(f"cannot open {path} even with shared access (Win32 error {ctypes.GetLastError()})")
    with os.fdopen(msvcrt.open_osfhandle(h, os.O_RDONLY | os.O_BINARY), "rb") as f:
        return f.read()


def load_workbook_safely(path: Path) -> openpyxl.Workbook:
    """openpyxl gets PermissionError while the workbook is open in Excel; fall back to a shared read."""
    try:
        return openpyxl.load_workbook(path, read_only=True)
    except PermissionError:
        if sys.platform != "win32":
            raise
        return openpyxl.load_workbook(io.BytesIO(_read_bytes_shared(path)), read_only=True)


def load_osca(path: Path) -> dict[str, dict]:
    """Return {unit_group_code: {title, lead, skill_level, occupations:[{code,title,alt,spec}]}}."""
    wb = load_workbook_safely(path)
    t2 = [r for r in wb["Table 2"].iter_rows(values_only=True) if r[0] and str(r[0]).strip().isdigit()]
    t1 = [r for r in wb["Table 1"].iter_rows(values_only=True) if r[0] and str(r[0]).strip().isdigit()]

    groups: dict[str, dict] = {}
    for r in t2:
        code = str(r[0]).strip()
        if len(code) != 4:
            continue
        groups[code] = {
            "code": code,
            "title": str(r[1]).strip(),
            "lead": str(r[2] or "").strip(),
            "skill_level": str(r[4] or "").strip(),
            "occupations": [],
        }
    for r in t1:
        code = str(r[0]).strip()
        ug = code[:4]
        if ug in groups:
            groups[ug]["occupations"].append({
                "code": code,
                "title": str(r[1] or "").strip(),
                "alt": str(r[2] or "").strip(),
                "spec": str(r[9] or "").strip(),
            })
    return groups


def format_group(g: dict) -> str:
    occ_titles = "; ".join(o["title"] for o in g["occupations"] if o["title"])
    alts = "; ".join(x for o in g["occupations"] for x in (o["alt"], o["spec"]) if x)
    line = f"{g['code']} {g['title']} (skill level {g['skill_level']}): {g['lead']}"
    if occ_titles:
        line += f" Occupations: {occ_titles}."
    if alts:
        line += f" Also known as: {alts}."
    return line


def build_system_prompt(groups: dict[str, dict]) -> list[dict]:
    instructions = """You are an expert occupation coder for the Australian Bureau of Statistics' OSCA \
(Occupation Standard Classification for Australia, 2024). You classify job advertisements to the \
4-digit OSCA unit group that best matches the work actually performed.

Coding rules:
- Code the occupation, not the industry or the employer. A software engineer at a bank is still a software engineer.
- Code the substantive work described in the responsibilities, not the title alone. Titles are inflated and vary by country: \
"Account Executive" in a US tech company is usually a sales representative, not a manager; a "Manager" who has no reports and \
does hands-on work is coded to the hands-on unit group; "Head of" / "Director" / "VP" with genuine control over a function is a manager.
- OSCA skill levels: 1 = bachelor degree or higher / 5+ years experience; 2 = diploma; 3 = certificate III/IV or trade; \
4 = certificate II; 5 = certificate I or short on-the-job training. Use the ad's stated requirements to judge the ad's skill level \
and let that inform the choice between adjacent unit groups.
- Most ads in this dataset come from technology, startup and knowledge-work employers in many countries. Map foreign titles to \
the closest Australian occupation.
- Give up to 3 candidates ordered best first. Confidence is your probability that the candidate is the correct unit group; \
the confidences do not need to sum to 1, but should sum to at most 1.
- Set issue when the ad is not a single, real, identifiable job: talent-community / expression-of-interest pages, ads covering \
several unrelated roles, placeholder or test postings, ads not in English (still attempt a code if you can read it), \
or ads with too little text to classify.

The full OSCA unit group list follows. Return codes only from this list."""
    listing = "\n".join(format_group(g) for g in sorted(groups.values(), key=lambda g: g["code"]))
    return [
        {"type": "text", "text": instructions},
        {"type": "text", "text": "OSCA UNIT GROUPS\n" + listing, "cache_control": {"type": "ephemeral"}},
    ]


# --------------------------------------------------------------------------- output schema
class Candidate(BaseModel):
    code: str = Field(description="4-digit OSCA unit group code")
    title: str = Field(description="unit group title exactly as listed")
    confidence: float = Field(description="probability between 0 and 1 that this is the correct unit group")
    rationale: str = Field(description="one sentence: which duties / requirements drove this choice")


class Label(BaseModel):
    issue: Literal["none", "talent_community", "multiple_roles", "not_a_job", "non_english", "insufficient_text"]
    ad_skill_level: int = Field(description="OSCA skill level 1-5 implied by the ad's stated requirements")
    candidates: list[Candidate] = Field(description="best first, 1 to 3 items; may be empty only if issue is not_a_job or insufficient_text")


# --------------------------------------------------------------------------- per-ad call
def ad_to_message(ad: dict) -> str:
    text = (ad.get("description_text") or "").strip()
    if len(text) > MAX_AD_CHARS:
        text = text[:MAX_AD_CHARS] + "\n[... truncated ...]"
    locs = ", ".join(ad.get("locations") or []) or "(none given)"
    return (
        f"JOB TITLE: {ad.get('title')}\n"
        f"DEPARTMENT: {ad.get('department') or '(none)'}\n"
        f"EMPLOYMENT TYPE: {ad.get('employment_type') or '(none)'}\n"
        f"LOCATION: {locs}\n"
        f"EMPLOYER SLUG: {ad.get('source_slug')}\n\n"
        f"AD TEXT:\n{text}\n\n"
        "Classify this ad to OSCA unit groups."
    )


def label_one(client: anthropic.Anthropic, system: list[dict], ad: dict, model: dict, effort: str) -> dict:
    extra = {"output_config": {"effort": effort}} if model["effort"] else {}
    resp = client.messages.parse(
        model=model["id"],
        max_tokens=2000,
        system=system,
        messages=[{"role": "user", "content": ad_to_message(ad)}],
        output_format=Label,
        **extra,
    )
    if resp.stop_reason == "refusal":
        return {"id": ad["id"], "error": "refusal", "stop_details": str(resp.stop_details)}
    parsed: Label = resp.parsed_output
    u = resp.usage
    return {
        "id": ad["id"],
        "label": parsed.model_dump(),
        "usage": {
            "input": u.input_tokens,
            "cache_write": u.cache_creation_input_tokens or 0,
            "cache_read": u.cache_read_input_tokens or 0,
            "output": u.output_tokens,
        },
        "model": resp.model,
    }


# --------------------------------------------------------------------------- review sheet
def write_review_xlsx(ads: list[dict], results: dict[str, dict], groups: dict[str, dict], path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "review"
    header = [
        "id", "title", "company", "department", "location", "url",
        "issue", "ad_skill_level",
        "code_1", "title_1", "conf_1", "rationale_1", "ug_skill_1", "skill_mismatch_1",
        "code_2", "title_2", "conf_2", "rationale_2",
        "code_3", "title_3", "conf_3", "rationale_3",
        "gold_code", "gold_title", "reviewer_notes", "error",
    ]
    ws.append(header)
    for ad in ads:
        r = results.get(ad["id"])
        row = [
            ad["id"], ad.get("title"), ad.get("source_slug"), ad.get("department"),
            ", ".join(ad.get("locations") or []), ad.get("apply_url"),
        ]
        if not r or "label" not in r:
            row += [None] * 16 + [None, None, None, (r or {}).get("error", "not labelled")]
            ws.append(row)
            continue
        lab = r["label"]
        row += [lab["issue"], lab["ad_skill_level"]]
        cands = lab["candidates"][:3]
        for i in range(3):
            if i < len(cands):
                c = cands[i]
                valid = c["code"] in groups
                title = groups[c["code"]]["title"] if valid else f"INVALID CODE: {c['title']}"
                row += [c["code"], title, round(c["confidence"], 2), c["rationale"]]
                if i == 0:
                    ug_sl = groups[c["code"]]["skill_level"] if valid else ""
                    mismatch = valid and ug_sl and str(lab["ad_skill_level"]) not in ug_sl.replace(" ", "").split(";")
                    row += [ug_sl, "Y" if mismatch else ""]
            else:
                row += [None, None, None, None] + ([None, None] if i == 0 else [])
        row += [None, None, None, None]
        ws.append(row)

    # a second sheet with the code list so reviewers can look codes up with a formula or filter
    ws2 = wb.create_sheet("osca_unit_groups")
    ws2.append(["code", "title", "skill_level", "lead_statement", "occupations"])
    for g in sorted(groups.values(), key=lambda g: g["code"]):
        ws2.append([g["code"], g["title"], g["skill_level"], g["lead"],
                    "; ".join(o["title"] for o in g["occupations"])])

    widths = {"B": 40, "C": 18, "E": 28, "J": 34, "L": 60, "P": 34, "R": 60, "T": 34, "V": 60, "X": 30, "Y": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# --------------------------------------------------------------------------- main
def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input", help="JSONL of ads from sample_jobs.py")
    ap.add_argument("--osca", default=str(DEFAULT_OSCA))
    ap.add_argument("--tag", default=None, help="output name suffix; default = input file stem")
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--model", default=DEFAULT_MODEL, choices=list(MODELS), help="short model name; output files carry it")
    ap.add_argument("--effort", default="medium", choices=["low", "medium", "high", "xhigh", "max"],
                    help="ignored for models that do not support effort")
    ap.add_argument("--dry-run", action="store_true", help="build prompts and report sizes; no API calls")
    ap.add_argument("--fresh", action="store_true", help="discard existing output for this input+model and relabel everything")
    args = ap.parse_args()

    groups = load_osca(Path(args.osca))
    system = build_system_prompt(groups)
    ads = [json.loads(l) for l in open(args.input, encoding="utf-8")]
    if args.limit:
        ads = ads[: args.limit]

    model = MODELS[args.model]
    PRICE = model["price"]
    tag = (args.tag or Path(args.input).stem.replace("sample_", "")) + "_" + args.model
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    jsonl_path = OUT_DIR / f"osca_labels_{tag}.jsonl"
    xlsx_path = OUT_DIR / f"osca_review_{tag}.xlsx"

    n_occ = sum(len(g["occupations"]) for g in groups.values())
    print(f"OSCA: {len(groups)} unit groups, {n_occ} occupations; system prompt ~{sum(len(b['text']) for b in system)//4:,} tokens")
    print(f"ads: {len(ads)} from {args.input}; model {model['id']}")

    if args.dry_run:
        print("\n--- first ad message ---\n" + ad_to_message(ads[0])[:1500] + "\n...")
        sys_tokens = sum(len(b["text"]) for b in system) // 4
        est_in = len(ads) * (len(ad_to_message(ads[0])) // 4 + 300)
        est_out = len(ads) * 300
        cost = (sys_tokens * PRICE["cache_write"] + len(ads) * sys_tokens * PRICE["cache_read"]
                + est_in * PRICE["input"] + est_out * PRICE["output"]) / 1e6
        print(f"\nestimated cost at {model['id']} list price: ~${cost:.2f} "
              f"(one cache write, {len(ads)} cache reads of ~{sys_tokens:,} tokens, "
              f"~{est_in:,} uncached input tokens, ~{est_out:,} output tokens)")
        return

    load_env()  # entry-level-work/.env overrides any ANTHROPIC_API_KEY already in the environment
    if not os.environ.get("ANTHROPIC_API_KEY"):
        sys.exit("ANTHROPIC_API_KEY is not set. Copy .env.example to .env and add your key, or set the variable.")
    client = anthropic.Anthropic(max_retries=4)

    # resume
    results: dict[str, dict] = {}
    if args.fresh and jsonl_path.exists():
        jsonl_path.unlink()
        print("--fresh: discarded previous labels for this input and model")
    if jsonl_path.exists():
        for line in open(jsonl_path, encoding="utf-8"):
            r = json.loads(line)
            if "label" in r:
                results[r["id"]] = r
        print(f"resuming: {len(results)} already labelled")
    todo = [a for a in ads if a["id"] not in results]

    lock = threading.Lock()
    out = open(jsonl_path, "a", encoding="utf-8")
    t0 = time.time()

    # first call alone so the cache is written once before the workers fan out
    if todo:
        first = todo.pop(0)
        r = label_one(client, system, first, model, args.effort)
        results[first["id"]] = r
        out.write(json.dumps(r, ensure_ascii=False) + "\n")
        out.flush()
        print(f"[1/{len(ads)}] {first['title'][:60]!r} -> "
              f"{[c['code'] for c in r.get('label', {}).get('candidates', [])]}  cache_write={r.get('usage', {}).get('cache_write')}")

    def work(ad: dict) -> dict:
        try:
            return label_one(client, system, ad, model, args.effort)
        except anthropic.APIError as e:
            return {"id": ad["id"], "error": f"{type(e).__name__}: {e}"}

    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = {ex.submit(work, ad): ad for ad in todo}
        for i, fut in enumerate(as_completed(futs), start=len(results) + 1):
            ad = futs[fut]
            r = fut.result()
            with lock:
                results[ad["id"]] = r
                out.write(json.dumps(r, ensure_ascii=False) + "\n")
                out.flush()
            codes = [c["code"] for c in r.get("label", {}).get("candidates", [])] if "label" in r else r.get("error")
            print(f"[{i}/{len(ads)}] {ad['title'][:60]!r} -> {codes}")
    out.close()

    # summary
    ok = [r for r in results.values() if "label" in r]
    errs = [r for r in results.values() if "label" not in r]
    usage = {k: sum(r["usage"][k] for r in ok) for k in PRICE}
    cost = sum(usage[k] * PRICE[k] / 1e6 for k in PRICE)
    invalid = sum(1 for r in ok for c in r["label"]["candidates"] if c["code"] not in groups)
    issues = {}
    for r in ok:
        issues[r["label"]["issue"]] = issues.get(r["label"]["issue"], 0) + 1
    print(f"\nlabelled {len(ok)}, errors {len(errs)}, invalid codes returned {invalid}, {time.time()-t0:.0f}s")
    print(f"tokens: {usage}  cost ~${cost:.2f}")
    print(f"issues: {issues}")

    write_review_xlsx(ads, results, groups, xlsx_path)
    print(f"wrote {jsonl_path}\nwrote {xlsx_path}")


if __name__ == "__main__":
    main()
