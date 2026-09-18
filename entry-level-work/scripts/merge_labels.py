"""Merge OSCA occupation labels and seniority labels onto the ads, one row per ad.

Joins on the ad id. For each label the gold value from the review workbook is used when a
reviewer has filled it in, otherwise the model's top choice. A `*_source` column records
which one was used. Writes CSV and Excel plus a unit-group x level crosstab.

Usage:
    python scripts/merge_labels.py data/samples/sample_100_australia.jsonl                 # default model
    python scripts/merge_labels.py data/samples/sample_100_australia.jsonl --model haiku-4-5
    python scripts/merge_labels.py data/samples/sample_100.jsonl --osca-model sonnet-5 --seniority-model haiku-4-5

Output: data/merged/merged_<tag>.csv, merged_<tag>.xlsx (sheets: merged, crosstab).
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from label_osca import DEFAULT_MODEL, DEFAULT_OSCA, MODELS, load_osca  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
LABEL_DIR = ROOT / "data" / "labels"
OUT_DIR = ROOT / "data" / "merged"


def read_jsonl(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    out = {}
    for line in open(path, encoding="utf-8"):
        r = json.loads(line)
        if "label" in r:
            out[r["id"]] = r["label"]
    return out


def read_gold(path: Path, gold_col: str, notes_col: str = "reviewer_notes") -> dict[str, tuple[str, str]]:
    """{id: (gold value, reviewer notes)} for rows where the reviewer filled the gold column."""
    if not path.exists():
        return {}
    ws = openpyxl.load_workbook(path, read_only=True)["review"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    i_id, i_gold = header.index("id"), header.index(gold_col)
    i_notes = header.index(notes_col) if notes_col in header else None
    gold = {}
    for r in rows:
        if r[i_id] and r[i_gold] not in (None, ""):
            gold[str(r[i_id])] = (str(r[i_gold]).strip(), str(r[i_notes]).strip() if i_notes is not None and r[i_notes] else "")
    return gold


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input", help="JSONL of ads from sample_jobs.py")
    ap.add_argument("--model", default=DEFAULT_MODEL, choices=list(MODELS), help="model used for both label sets")
    ap.add_argument("--osca-model", default=None, choices=list(MODELS), help="override model for OSCA labels")
    ap.add_argument("--seniority-model", default=None, choices=list(MODELS), help="override model for seniority labels")
    ap.add_argument("--osca", default=str(DEFAULT_OSCA))
    ap.add_argument("--tag", default=None)
    args = ap.parse_args()

    tag = args.tag or Path(args.input).stem.replace("sample_", "")
    osca_model = args.osca_model or args.model
    sen_model = args.seniority_model or args.model

    ads = [json.loads(l) for l in open(args.input, encoding="utf-8")]
    groups = load_osca(Path(args.osca))
    osca = read_jsonl(LABEL_DIR / f"osca_labels_{tag}_{osca_model}.jsonl")
    osca_gold = read_gold(LABEL_DIR / f"osca_review_{tag}_{osca_model}.xlsx", "gold_code")
    sen = read_jsonl(LABEL_DIR / f"seniority_labels_{tag}_{sen_model}.jsonl")
    sen_gold = read_gold(LABEL_DIR / f"seniority_review_{tag}_{sen_model}.xlsx", "gold_level")
    print(f"ads {len(ads)} | osca labels {len(osca)} (gold {len(osca_gold)}) [{osca_model}] "
          f"| seniority labels {len(sen)} (gold {len(sen_gold)}) [{sen_model}]")
    if not osca and not sen:
        sys.exit("no label files found for this input/model; run label_osca.py and label_seniority.py first")

    header = ["id", "title", "company", "department", "employment_type", "location", "remote", "posted_at", "url",
              "osca_code", "osca_title", "osca_major_group", "osca_skill_level", "osca_source", "osca_confidence",
              "osca_issue", "osca_notes",
              "level", "level_source", "is_internship", "level_confidence", "level_evidence", "level_notes"]
    rows = []
    for ad in ads:
        r = {k: None for k in header}
        r.update(id=ad["id"], title=ad.get("title"), company=ad.get("source_slug"), department=ad.get("department"),
                 employment_type=ad.get("employment_type"), location=", ".join(ad.get("locations") or []),
                 remote=ad.get("remote"), posted_at=ad.get("posted_at"), url=ad.get("apply_url"))

        o = osca.get(ad["id"])
        if ad["id"] in osca_gold:
            r["osca_code"], r["osca_notes"] = osca_gold[ad["id"]]
            r["osca_source"] = "gold"
        elif o and o["candidates"]:
            r["osca_code"] = o["candidates"][0]["code"]
            r["osca_source"] = "model"
            r["osca_confidence"] = round(o["candidates"][0]["confidence"], 2)
        if o:
            r["osca_issue"] = o["issue"] if o["issue"] != "none" else None
        g = groups.get(r["osca_code"] or "")
        if g:
            r["osca_title"], r["osca_major_group"], r["osca_skill_level"] = g["title"], g["code"][0], g["skill_level"]
        elif r["osca_code"]:
            r["osca_title"] = "UNKNOWN CODE"

        s = sen.get(ad["id"])
        if ad["id"] in sen_gold:
            r["level"], r["level_notes"] = sen_gold[ad["id"]]
            r["level_source"] = "gold"
        elif s:
            r["level"], r["level_source"] = s["level"], "model"
        if s:
            r["is_internship"] = s["is_internship"]
            r["level_confidence"] = round(s["confidence"], 2)
            r["level_evidence"] = s["evidence"]
        rows.append(r)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    csv_path, xlsx_path = OUT_DIR / f"merged_{tag}.csv", OUT_DIR / f"merged_{tag}.xlsx"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=header)
        w.writeheader()
        w.writerows(rows)

    # crosstab: unit group x level
    xt: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for r in rows:
        if r["osca_code"]:
            xt[(r["osca_code"], r["osca_title"])][r["level"] or "unlabelled"] += 1
    levels = ["entry_level", "experienced", "unclear", "unlabelled"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "merged"
    ws.append(header)
    for r in rows:
        ws.append([r[k] for k in header])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws2 = wb.create_sheet("crosstab")
    ws2.append(["osca_code", "osca_title", *levels, "total"])
    for (code, title), c in sorted(xt.items(), key=lambda kv: -sum(kv[1].values())):
        ws2.append([code, title, *(c[l] for l in levels), sum(c.values())])
    wb.save(xlsx_path)

    lv = Counter(r["level"] or "unlabelled" for r in rows)
    print(f"levels: {dict(lv)} | unit groups: {len(xt)} | ads with both labels: "
          f"{sum(1 for r in rows if r['osca_code'] and r['level'])}")
    print(f"wrote {csv_path}\nwrote {xlsx_path}")


if __name__ == "__main__":
    main()
